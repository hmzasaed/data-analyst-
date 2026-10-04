"""Readers for delimited text and XLSX workbooks.

Both readers produce the same thing: a `NormalizedTable` of strings plus the
reader's own evidence about what it saw (encoding, delimiter, sheet, skipped
rows). Text is kept as text at this stage; type coercion is a separate,
inspectable cleaning step so that "why is this column a number" is answerable.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import polars as pl
from insightflow_contracts import IngestionError, UnsupportedInputError
from openpyxl import load_workbook

from insightflow_ingestion.normalize import (
    NormalizedTable,
    normalise_column_names,
    normalise_whitespace,
)
from insightflow_ingestion.sniffing import (
    detect_delimiter,
    detect_encoding,
    detect_format,
)

#: Guard rails for degenerate sheets.
MAX_PREVIEW_ROWS = 1_000


@dataclass
class ReadOptions:
    sheet_name: str | None = None
    delimiter: str | None = None
    encoding: str | None = None
    has_header: bool = True
    max_rows: int | None = None
    max_columns: int | None = None
    #: Treat these as missing on read (sheets use them to mean "no value").
    na_values: tuple[str, ...] = (
        "",
        "na",
        "n/a",
        "N/A",
        "null",
        "NULL",
        "None",
        "none",
        "nan",
        "NaN",
        "-",
        "--",
        "#N/A",
        "#NULL!",
        "#VALUE!",
        "#DIV/0!",
    )


@dataclass
class ReadResult:
    table: NormalizedTable
    sheet_name: str | None = None
    available_sheets: list[str] = field(default_factory=list)
    encoding: str | None = None
    delimiter: str | None = None
    detected_format: str = "csv"
    skipped_rows: int = 0
    notes: list[str] = field(default_factory=list)
    read_seconds: float = 0.0

    def evidence(self) -> dict[str, Any]:
        return {
            "sheet_name": self.sheet_name,
            "available_sheets": self.available_sheets,
            "encoding": self.encoding,
            "delimiter": self.delimiter,
            "detected_format": self.detected_format,
            "skipped_rows": self.skipped_rows,
            "notes": self.notes,
        }


def list_sheets(path: str | Path) -> list[str]:
    """Return worksheet names in workbook order, without loading cell data."""
    file_path = Path(path)
    if detect_format(file_path) != "xlsx":
        return []
    try:
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:  # pragma: no cover - defensive
        raise IngestionError("workbook could not be opened", path=str(file_path), cause=type(exc).__name__) from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _read_csv(path: Path, options: ReadOptions) -> ReadResult:
    encoding = options.encoding or detect_encoding(path).encoding
    delimiter = options.delimiter
    notes: list[str] = []
    if delimiter is None:
        sniff = detect_delimiter(path, encoding=encoding)
        delimiter = sniff.delimiter
        notes.append(f"delimiter: {sniff.rationale}")
        if sniff.delimiter is None:
            notes.append("no delimiter detected; treating the file as a single column")

    read_options: dict[str, Any] = {
        "has_header": options.has_header,
        "infer_schema_length": 0,  # keep everything as text; coerce explicitly later
        "null_values": list(options.na_values),
        "ignore_errors": True,
        "truncate_ragged_lines": True,
        "low_memory": True,
    }
    if delimiter:
        read_options["separator"] = delimiter

    try:
        frame = pl.read_csv(path, encoding=encoding, **read_options)
    except Exception as exc:
        raise IngestionError(
            "delimited text could not be parsed",
            path=str(path),
            encoding=encoding,
            delimiter=delimiter,
            cause=type(exc).__name__,
        ) from exc

    if options.max_columns and frame.width > options.max_columns:
        notes.append(f"truncated to the first {options.max_columns} columns")
        frame = frame.select(list(frame.columns)[: options.max_columns])
    if options.max_rows is not None and frame.height > options.max_rows:
        notes.append(f"row limit {options.max_rows} reached during ingestion")
        frame = frame.head(options.max_rows)

    table = _finalise(
        frame,
        original_header=_raw_header(path, encoding, delimiter),
        notes=notes,
    )
    return ReadResult(
        table=table,
        sheet_name=None,
        available_sheets=[],
        encoding=encoding,
        delimiter=delimiter,
        detected_format="csv",
        notes=notes,
    )


def _read_xlsx(path: Path, options: ReadOptions) -> ReadResult:
    available = list_sheets(path)
    if not available:
        raise UnsupportedInputError("workbook contains no worksheets", path=str(path))

    sheet_name = options.sheet_name
    if sheet_name is None:
        sheet_name = available[0]
    if sheet_name not in available:
        raise UnsupportedInputError(
            f"sheet '{sheet_name}' not found",
            available_sheets=available,
            suggestions=[f"Choose one of: {', '.join(available)}"],
        )

    notes: list[str] = []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise IngestionError("workbook could not be opened", path=str(path)) from exc

    try:
        worksheet = workbook[sheet_name]
        # data_only=True means formula cells yield their cached result and the
        # formula text is never evaluated or interpreted.
        rows_iter = worksheet.iter_rows(values_only=True)
        header: list[Any] | None = None
        if options.has_header:
            header = next(rows_iter, None)

        collected: list[tuple[Any, ...]] = []
        skipped = 0
        for row in rows_iter:
            if row is None or all(cell is None for cell in row):
                skipped += 1
                continue
            if options.max_rows is not None and len(collected) >= options.max_rows:
                skipped += 1
                continue
            collected.append(tuple(row))
            if len(collected) >= MAX_PREVIEW_ROWS * 100:
                break
    finally:
        workbook.close()

    if header is None and not collected:
        raise UnsupportedInputError(
            f"sheet '{sheet_name}' is empty",
            path=str(path),
            sheet=sheet_name,
            suggestions=["Pick a sheet that contains data, or upload a different file."],
        )

    width = max(
        len(header) if header else 0,
        max((len(row) for row in collected), default=0),
    )
    if options.max_columns and width > options.max_columns:
        notes.append(f"truncated to the first {options.max_columns} columns")
        width = options.max_columns

    def _pad(row: tuple[Any, ...]) -> list[str | None]:
        cells = list(row[:width])
        cells.extend([None] * (width - len(cells)))
        return [None if cell is None else str(cell) for cell in cells]

    if header is None:
        header = [f"column_{i + 1}" for i in range(width)]
    header = [str(cell) if cell is not None else "" for cell in header[:width]]

    data = {header[i]: [_pad(row)[i] for row in collected] for i in range(width)}
    frame = pl.DataFrame(data, strict=False)
    table = _finalise(frame, original_header=header, notes=notes)
    return ReadResult(
        table=table,
        sheet_name=sheet_name,
        available_sheets=available,
        detected_format="xlsx",
        skipped_rows=skipped,
        notes=notes,
    )


def _raw_header(path: Path, encoding: str, delimiter: str | None) -> list[str]:
    """Original (unnormalised) header names, for rename reporting.

    Polars silently de-duplicates repeated headers on read, so the raw line is
    needed to be able to tell the user their duplicates were renamed.
    """
    if not delimiter:
        return []
    try:
        with path.open("r", encoding=encoding, errors="replace", newline="") as handle:
            first = handle.readline()
    except OSError:  # pragma: no cover - defensive
        return []
    import csv as _csv

    try:
        return next(_csv.reader([first], delimiter=delimiter))
    except (StopIteration, _csv.Error):
        return []


def _finalise(
    frame: pl.DataFrame,
    *,
    original_header: list[str] | None = None,
    notes: list[str],
) -> NormalizedTable:
    """Apply whitespace and column-name normalisation to a raw frame."""
    for column in frame.columns:
        if frame.schema[column] == pl.String:
            frame = frame.with_columns(pl.col(column).str.strip_chars().alias(column))

    current = list(frame.columns)
    # Prefer the raw header: it is the only place duplicate or empty names survive.
    source_names = (
        [str(name) for name in original_header]
        if original_header is not None and len(original_header) == frame.width
        else current
    )
    stripped = [normalise_whitespace(name) or name for name in source_names]
    resolved, collisions = normalise_column_names(stripped)

    for collision in collisions:
        notes.append(
            f"column {collision.position + 1} ('{collision.original}') -> "
            f"'{collision.normalised}' ({collision.reason})"
        )

    rename_map = {
        current[index]: resolved[index]
        for index in range(len(current))
        if current[index] != resolved[index]
    }
    if rename_map:
        frame = frame.rename(rename_map)

    if frame.width == 0:
        raise IngestionError("table has no columns")

    return NormalizedTable(
        frame=frame,
        original_columns=[str(name) for name in source_names],
        notes=notes,
    )


def read_table(path: str | Path, options: ReadOptions | None = None) -> ReadResult:
    """Read any supported file into a normalised string table.

    Raises `UnsupportedInputError` for content the platform will not guess about,
    and `IngestionError` for content it cannot parse.
    """
    read_options = options or ReadOptions()
    file_path = Path(path)
    detected = detect_format(file_path)
    if detected == "xlsx":
        return _read_xlsx(file_path, read_options)
    if detected == "csv":
        return _read_csv(file_path, read_options)
    raise UnsupportedInputError(
        f"unsupported file content detected as '{detected}'",
        path=str(file_path),
        suggestions=["Provide a .csv, .tsv or .xlsx file."],
    )


def preview_rows(table: NormalizedTable, limit: int = 20) -> list[dict[str, Any]]:
    """Small, safe preview for the UI. Values are strings, capped in length."""
    safe_limit = max(0, min(limit, MAX_PREVIEW_ROWS))
    sliced = table.frame.head(safe_limit)
    return [
        {column: (str(row[column])[:200] if row[column] is not None else None) for column in sliced.columns}
        for row in sliced.iter_rows(named=True)
    ]
