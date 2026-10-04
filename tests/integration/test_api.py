"""API contract: upload, review, findings, exports, tenant isolation."""

from __future__ import annotations

import pytest
from fastapi.testclient import TestClient

pytestmark = pytest.mark.integration


@pytest.fixture
def client(tmp_path, monkeypatch, sales_csv):
    from insightflow_contracts import reset_settings

    monkeypatch.setenv("INSIGHTFLOW_ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    monkeypatch.setenv(
        "INSIGHTFLOW_STATE_DATABASE_URL", f"sqlite:///{(tmp_path / 'state.db').as_posix()}"
    )
    monkeypatch.setenv("INSIGHTFLOW_LLM_ENABLED", "false")
    monkeypatch.setenv("INSIGHTFLOW_INLINE_EXECUTION", "true")
    reset_settings()

    import insightflow_api.service as app_module

    app_module.state = _fresh_state(tmp_path)
    app_module.artifacts = _fresh_store(tmp_path)
    app_module.factory = _fresh_factory(tmp_path, app_module)
    with TestClient(app_module.app) as test_client:
        yield test_client


def _fresh_state(tmp_path):
    from insightflow_storage import StateStore

    return StateStore(f"sqlite:///{(tmp_path / 'api.db').as_posix()}")


def _fresh_store(tmp_path):
    from insightflow_storage import ArtifactStore

    return ArtifactStore(tmp_path / "api-artifacts")


def _fresh_factory(tmp_path, app_module):
    from insightflow_ingestion import IngestionService
    from insightflow_orchestration import JobFactory

    settings = app_module.settings
    store = _fresh_store(tmp_path)
    return JobFactory(
        state=app_module.state,
        artifact_store=store,
        ingestion=IngestionService(settings.artifact_path("raw")),
    )


class TestHealthAndCatalogue:
    def test_health_reports_the_runtime(self, client):
        payload = client.get("/health").json()
        assert payload["status"] == "ok"
        assert payload["llm_enabled"] is False

    def test_tool_catalogue_is_exposed(self, client):
        payload = client.get("/api/v1/tools").json()
        assert payload["count"] >= 15
        names = {tool["name"] for tool in payload["tools"]}
        assert "compare_groups" in names
        for tool in payload["tools"]:
            assert tool["summary"]
            assert tool["parameters"]


class TestAuthentication:
    def test_requests_without_an_owner_are_refused(self, client, sales_csv):
        response = client.post(
            "/api/v1/jobs", files={"file": ("sales.csv", sales_csv.read_bytes(), "text/csv")}
        )
        assert response.status_code == 401

    def test_a_missing_job_is_not_found(self, client):
        assert client.get("/api/v1/jobs/nope", headers={"X-Owner-Id": "alice"}).status_code == 404


class TestUploadAndReport:
    def _upload(self, client, path, *, question="", owner="alice"):
        return client.post(
            "/api/v1/jobs",
            headers={"X-Owner-Id": owner},
            files={"file": (path.name, path.read_bytes(), "text/csv")},
            data={"question": question} if question else {},
        )

    def test_upload_returns_a_completed_job(self, client, sales_csv):
        response = self._upload(client, sales_csv, question="How did revenue trend?")
        assert response.status_code == 201
        payload = response.json()
        assert payload["file_hash"]
        assert payload["status"] == "completed", payload

    def test_full_flow_produces_a_report(self, client, sales_csv):
        job_id = self._upload(client, sales_csv, question="Compare revenue by region.").json()["job_id"]
        headers = {"X-Owner-Id": "alice"}

        progress = client.get(f"/api/v1/jobs/{job_id}", headers=headers).json()
        assert progress["progress"] >= 0

        findings = client.get(f"/api/v1/jobs/{job_id}/findings", headers=headers).json()
        assert findings["findings"]
        assert findings["verification"]

        report = client.get(f"/api/v1/jobs/{job_id}/report", headers=headers)
        assert report.status_code == 200
        assert "text/html" in report.headers["content-type"]
        assert "Executive summary" in report.text

    def test_markdown_and_json_exports(self, client, sales_csv):
        job_id = self._upload(client, sales_csv, question="overview").json()["job_id"]
        headers = {"X-Owner-Id": "alice"}

        markdown = client.get(
            f"/api/v1/jobs/{job_id}/report", headers=headers, params={"format": "markdown"}
        )
        assert markdown.status_code == 200
        assert markdown.text.startswith("#")

        payload = client.get(
            f"/api/v1/jobs/{job_id}/report", headers=headers, params={"format": "json"}
        ).json()
        assert payload["document"]["sections"]
        assert payload["findings"]

    def test_profile_endpoint_returns_the_profile(self, client, sales_csv):
        job_id = self._upload(client, sales_csv).json()["job_id"]
        profile = client.get(
            f"/api/v1/jobs/{job_id}/profile", headers={"X-Owner-Id": "alice"}
        ).json()
        assert profile["row_count"] == 240
        assert profile["column_count"] == 7
        assert any(column["inferred_type"] == "number" for column in profile["columns"])

    def test_cleaning_proposals_are_reviewable(self, client, sales_csv):
        job_id = self._upload(client, sales_csv).json()["job_id"]
        headers = {"X-Owner-Id": "alice"}
        cleaning = client.get(f"/api/v1/jobs/{job_id}/cleaning", headers=headers).json()
        assert cleaning["summary"]["total_operations"] >= 0
        for operation in cleaning["operations"]:
            assert operation["reason"]

        if cleaning["operations"]:
            first = cleaning["operations"][0]["operation_id"]
            decision = client.post(
                f"/api/v1/jobs/{job_id}/cleaning/decision",
                headers=headers,
                json={"accept": [first]},
            )
            assert decision.status_code == 200
            assert first in decision.json()["approved"]

    def test_unknown_operation_id_is_refused(self, client, sales_csv):
        job_id = self._upload(client, sales_csv).json()["job_id"]
        response = client.post(
            f"/api/v1/jobs/{job_id}/cleaning/decision",
            headers={"X-Owner-Id": "alice"},
            json={"accept": ["op_does_not_exist"]},
        )
        assert response.status_code == 400

    def test_audit_trail_is_available(self, client, sales_csv):
        job_id = self._upload(client, sales_csv).json()["job_id"]
        audit = client.get(f"/api/v1/jobs/{job_id}/audit", headers={"X-Owner-Id": "alice"}).json()
        assert audit["events"]
        assert audit["stages"]

    def test_job_listing_is_scoped_to_the_owner(self, client, sales_csv):
        self._upload(client, sales_csv, owner="alice")
        assert client.get("/api/v1/jobs", headers={"X-Owner-Id": "alice"}).json()
        assert client.get("/api/v1/jobs", headers={"X-Owner-Id": "bob"}).json() == []

    def test_artifacts_are_served_only_to_the_owner(self, client, sales_csv):
        job_id = self._upload(client, sales_csv, question="compare by region").json()["job_id"]
        owner = {"X-Owner-Id": "alice"}
        report = client.get(f"/api/v1/jobs/{job_id}/report", headers=owner, params={"format": "json"})
        payload = report.json()
        artifact_ids = payload["artifacts"]["charts"] + payload["artifacts"]["tables"]
        assert artifact_ids, "the run produced no artifacts to serve"
        served = client.get(
            f"/api/v1/jobs/{job_id}/artifacts/{artifact_ids[0]['artifact_id']}", headers=owner
        )
        assert served.status_code == 200
        assert served.content
        # A different tenant cannot reach it.
        assert client.get(
            f"/api/v1/jobs/{job_id}/artifacts/{artifact_ids[0]['artifact_id']}",
            headers={"X-Owner-Id": "mallory"},
        ).status_code == 404

    def test_delete_removes_the_job(self, client, sales_csv):
        job_id = self._upload(client, sales_csv).json()["job_id"]
        owner = {"X-Owner-Id": "alice"}
        assert client.delete(f"/api/v1/jobs/{job_id}", headers=owner).json()["deleted"]
        assert client.get(f"/api/v1/jobs/{job_id}", headers=owner).status_code == 404


class TestUploadValidation:
    def test_unsupported_type_is_refused(self, client, tmp_path):
        path = tmp_path / "notes.exe"
        path.write_bytes(b"MZ\x00\x00")
        response = client.post(
            "/api/v1/jobs",
            headers={"X-Owner-Id": "alice"},
            files={"file": ("notes.exe", path.read_bytes(), "application/octet-stream")},
        )
        assert response.status_code == 415

    def test_empty_upload_is_refused(self, client, tmp_path):
        path = tmp_path / "empty.csv"
        path.write_bytes(b"")
        response = client.post(
            "/api/v1/jobs",
            headers={"X-Owner-Id": "alice"},
            files={"file": ("empty.csv", b"", "text/csv")},
        )
        assert response.status_code == 400

    def test_sheet_selection_is_validated(self, client, tmp_path):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.title = "Sales"
        workbook.active.append(["a"])
        path = tmp_path / "book.xlsx"
        workbook.save(path)
        response = client.post(
            "/api/v1/jobs",
            headers={"X-Owner-Id": "alice"},
            files={
                "file": (
                    "book.xlsx",
                    path.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"sheet_name": "Nonexistent"},
        )
        assert response.status_code == 415

    def test_xlsx_upload_is_accepted_and_sheets_listed(self, client, tmp_path):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.title = "Sales"
        workbook.active.append(["order_date", "revenue"])
        workbook.active.append(["2024-01-01", 100])
        workbook.active.append(["2024-02-01", 120])
        second = workbook.create_sheet("Notes")
        second.append(["note"])
        path = tmp_path / "book.xlsx"
        workbook.save(path)

        response = client.post(
            "/api/v1/jobs",
            headers={"X-Owner-Id": "alice"},
            files={
                "file": (
                    "book.xlsx",
                    path.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"question": "trend"},
        )
        assert response.status_code == 201
        assert "Sales" in response.json()["available_sheets"]

class TestReviewRoundTrip:
    def test_review_pauses_then_the_decision_completes_the_job(self, client, messy_csv):
        headers = {"X-Owner-Id": "alice"}
        upload = client.post(
            "/api/v1/jobs",
            headers=headers,
            files={"file": (messy_csv.name, messy_csv.read_bytes(), "text/csv")},
            data={"question": "Which region has the highest revenue?", "require_review": "true"},
        )
        assert upload.status_code == 201
        job_id = upload.json()["job_id"]
        assert upload.json()["status"] == "awaiting_review"

        progress = client.get(f"/api/v1/jobs/{job_id}", headers=headers).json()
        assert progress["require_review"] is True
        assert "cleaning_reviewed" not in progress["completed_stages"]
        assert progress["stages"]

        # Resuming without a decision is refused rather than silently skipping review.
        assert client.post(f"/api/v1/jobs/{job_id}/resume", headers=headers).status_code == 409

        cleaning = client.get(f"/api/v1/jobs/{job_id}/cleaning", headers=headers).json()
        assert cleaning["awaiting_review"] is True
        reviewable = [
            op["operation_id"] for op in cleaning["operations"] if op["review_level"] != "automatic"
        ]
        decision = client.post(
            f"/api/v1/jobs/{job_id}/cleaning/decision",
            headers=headers,
            json={"accept": reviewable[:1], "reject": reviewable[1:]},
        ).json()
        assert decision["status"] == "completed"
        assert decision["takes_effect"] is True

        after = client.get(f"/api/v1/jobs/{job_id}/cleaning", headers=headers).json()
        applied = {entry["operation_id"] for entry in after["applied"]}
        assert set(reviewable[:1]) <= applied
        assert not set(reviewable[1:]) & applied

    def test_embedded_report_has_no_header_dependent_images(self, client, sales_csv):
        headers = {"X-Owner-Id": "alice"}
        job_id = client.post(
            "/api/v1/jobs", headers=headers,
            files={"file": (sales_csv.name, sales_csv.read_bytes(), "text/csv")},
            data={"question": "How did revenue trend?"},
        ).json()["job_id"]
        html = client.get(f"/api/v1/jobs/{job_id}/report?embed=true", headers=headers).text
        assert "/artifacts/" not in html
        assert "data:image/png;base64" in html


class TestPreprocessingApi:
    def _upload(self, client, path, **data):
        return client.post(
            "/api/v1/jobs", headers={"X-Owner-Id": "alice"},
            files={"file": (path.name, path.read_bytes(), "text/csv")}, data=data,
        )

    def test_options_are_applied_and_the_export_is_downloadable(self, client, messy_csv):
        import json

        response = self._upload(
            client, messy_csv, question="Which region has the highest revenue?",
            preprocessing=json.dumps({"missing_numeric": "median", "scaling": "minmax"}),
        )
        assert response.status_code == 201, response.text
        job_id = response.json()["job_id"]
        result = client.get(f"/api/v1/jobs/{job_id}/preprocessing", headers={"X-Owner-Id": "alice"}).json()
        assert result["options"]["missing_numeric"] == "median"
        assert result["options"]["scaling"] == "minmax"
        assert result["export"]["artifact_id"]

        download = client.get(
            f"/api/v1/jobs/{job_id}/artifacts/{result['export']['artifact_id']}",
            headers={"X-Owner-Id": "alice"},
        )
        assert download.status_code == 200
        assert download.headers["content-type"].startswith("text/csv")
        header = download.text.splitlines()[0].split(",")
        assert len(header) == result["export"]["columns"]

    def test_invalid_options_are_refused(self, client, sales_csv):
        response = self._upload(client, sales_csv, preprocessing='{"missing_numeric": "guess"}')
        assert response.status_code == 422

    def test_the_report_is_organised_into_sections(self, client, sales_csv):
        job_id = self._upload(client, sales_csv, question="How did revenue trend?").json()["job_id"]
        html = client.get(f"/api/v1/jobs/{job_id}/report?embed=true", headers={"X-Owner-Id": "alice"}).text
        for anchor in ("executive_summary", "preprocessing", "descriptive_statistics",
                       "distributions", "key_findings", "appendix"):
            assert f'id="{anchor}"' in html
        assert "Additional charts" not in html


class TestFeaturesApi:
    def test_feature_agents_results_and_feature_set_download(self, client, sales_csv):
        import json

        headers = {"X-Owner-Id": "alice"}
        job_id = client.post(
            "/api/v1/jobs", headers=headers,
            files={"file": (sales_csv.name, sales_csv.read_bytes(), "text/csv")},
            data={"question": "What predicts revenue?", "features": json.dumps({"max_features": 4})},
        ).json()["job_id"]
        payload = client.get(f"/api/v1/jobs/{job_id}/features", headers=headers).json()
        engineering, selection = payload["engineering"], payload["selection"]
        assert engineering["target_column"] == "revenue"
        assert engineering["task"] == "regression"
        assert 0 < len(selection["selected"]) <= 4
        csv = client.get(f"/api/v1/jobs/{job_id}/artifacts/{selection['artifact_id']}", headers=headers).text
        assert csv.splitlines()[0].split(",") == [*selection["selected"], "revenue"]

        html = client.get(f"/api/v1/jobs/{job_id}/report?embed=true", headers=headers).text
        assert 'id="features"' in html and "Feature ranking" in html

    def test_invalid_feature_options_are_refused(self, client, sales_csv):
        response = client.post(
            "/api/v1/jobs", headers={"X-Owner-Id": "alice"},
            files={"file": (sales_csv.name, sales_csv.read_bytes(), "text/csv")},
            data={"features": '{"max_features": 0}'},
        )
        assert response.status_code == 422
